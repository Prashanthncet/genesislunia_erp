# -*- coding: utf-8 -*-
"""Generate the school-facing requirements template for Genesis Lumia ERP.

A fill-in workbook we send to a school before go-live. It collects everything
needed to configure the ERP: school profile, modules wanted, classes, subjects,
staff and roles, timetable, fees, admissions, hostel, attendance and leave,
exams and grading, notifications, existing data, reports, and sign-off.

Conventions inside the workbook:
  * light-blue cells  = the school fills these in
  * pale-gold row     = a worked example showing the expected format
  * white cells       = filled in by us (labels, guidance, pre-listed items)
  * "Start Here" tracks progress per section with live formulas.

Usage:  python scripts/generate_requirements_template.py
Output: Genesis_ERP_School_Requirements_Template.xlsx (repo root)
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "Genesis_ERP_School_Requirements_Template.xlsx"
RETURN_TO = "Jayanth  ·  jayanth.m@ncetmail.com"

# ── Theme (matches the ERP's navy design system) ─────────────────────────────
NAVY, DEEP, SECONDARY = "142B4A", "0B1F3A", "315A85"
PALE, PAGE, LINE, HAIR = "EAF2FB", "F5F7FB", "CBD5E1", "E2E8F0"
INK, MUTED, WHITE = "172B4D", "64748B", "FFFFFF"
GOLD_SOFT, GOLD_INK = "FBF4DE", "7A5C0E"
OK_SOFT, OK_INK = "ECFDF3", "166534"
WARN_SOFT, WARN_INK = "FEF3C7", "92400E"
SLATE_SOFT = "F1F5F9"
RED = "B91C1C"
FONT = "Arial"


def font(size=10, bold=False, color=INK, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


THIN = Side(style="thin", color=LINE)
HAIRLINE = Side(style="thin", color=HAIR)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ROW_RULE = Border(bottom=HAIRLINE)
TOP_LEFT = Alignment(vertical="top", wrap_text=True)
MID_LEFT = Alignment(vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

INPUT_FILL = fill(PALE)
EXAMPLE_FILL = fill(GOLD_SOFT)


# ── Shared dropdown lists ────────────────────────────────────────────────────
YES_NO = ["Yes", "No"]
YES_NO_LATER = ["Yes", "No", "Later"]
PRIORITY = ["High", "Medium", "Low"]


# ── Sheet helpers ────────────────────────────────────────────────────────────

class Sheet:
    """Builds one themed sheet top-to-bottom and records where its inputs live."""

    def __init__(self, wb: Workbook, name: str, title: str, subtitle: str, widths: list[float]):
        self.ws = wb.create_sheet(name)
        self.name = name
        self.ncols = len(widths)
        self.row = 1
        self.form_ranges: list[str] = []      # answer cells of question blocks
        self.table_keys: list[str] = []       # first input column of free tables
        self.checklists: list[tuple[str, int]] = []  # (answer range, item count)
        ws = self.ws
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = SECONDARY
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        last = get_column_letter(self.ncols)

        # Title band + subtitle band.
        ws.merge_cells(f"A1:{last}1")
        ws["A1"] = title
        ws["A1"].font = font(16, True, WHITE)
        ws["A1"].alignment = Alignment(vertical="center", indent=1)
        ws.merge_cells(f"A2:{last}2")
        ws["A2"] = subtitle
        ws["A2"].font = font(10, False, "D6E4F4")
        ws["A2"].alignment = Alignment(vertical="center", indent=1, wrap_text=True)
        for c in range(1, self.ncols + 1):
            ws.cell(1, c).fill = fill(NAVY)
            ws.cell(2, c).fill = fill(NAVY)
        ws.row_dimensions[1].height = 34
        ws.row_dimensions[2].height = 30

        # Back link + compact legend.
        ws["A3"] = '=HYPERLINK("#\'Start Here\'!A1","← Back to Start Here")'
        ws["A3"].font = font(9, True, SECONDARY, underline="single")
        ws.merge_cells("A3:B3")
        legend = ws.cell(3, 3)
        legend.value = "Light-blue cells: please fill in   ·   Gold row: example of the format   ·   * required"
        legend.font = font(9, False, MUTED, italic=True)
        ws.row_dimensions[3].height = 20
        self.row = 5

    # ---- blocks ----

    def heading(self, text: str, note: str | None = None):
        ws, r = self.ws, self.row
        last = get_column_letter(self.ncols)
        ws.merge_cells(f"A{r}:{last}{r}")
        ws.cell(r, 1).value = text.upper()
        ws.cell(r, 1).font = font(10, True, NAVY)
        ws.cell(r, 1).alignment = Alignment(vertical="bottom")
        for c in range(1, self.ncols + 1):
            ws.cell(r, c).border = Border(bottom=Side(style="medium", color=NAVY))
        ws.row_dimensions[r].height = 22
        self.row += 1
        if note:
            ws.merge_cells(f"A{self.row}:{last}{self.row}")
            ws.cell(self.row, 1).value = note
            ws.cell(self.row, 1).font = font(9, False, MUTED, italic=True)
            ws.cell(self.row, 1).alignment = MID_LEFT
            ws.row_dimensions[self.row].height = 18 if len(note) < 110 else 30
            self.row += 1

    def questions(self, items: list[tuple]):
        """Question block: # | Question | Your answer | How to answer.

        Each item: (question, guidance, required: bool, options: list|None).
        The answer cell spans columns C..(ncols-1); guidance sits in the last column.
        """
        ws = self.ws
        start = self.row
        ans_last = get_column_letter(self.ncols - 1)
        # On sheets whose column A is sized for a table, a narrow "#" column doesn't
        # exist — the numbered question spans A:B instead so it doesn't wrap.
        wide_a = (ws.column_dimensions["A"].width or 0) > 8
        for i, (q, guide, required, options) in enumerate(items, start=1):
            r = self.row
            label = f"{q} *" if required else q
            if wide_a:
                ws.merge_cells(f"A{r}:B{r}")
                ws.cell(r, 1, f"{i}.  {label}").font = font(10, required, INK)
                ws.cell(r, 1).alignment = MID_LEFT
            else:
                ws.cell(r, 1, i).font = font(9, False, MUTED)
                ws.cell(r, 1).alignment = CENTER
                ws.cell(r, 2, label).font = font(10, required, INK)
                ws.cell(r, 2).alignment = MID_LEFT
            if self.ncols - 1 > 3:
                ws.merge_cells(f"C{r}:{ans_last}{r}")
            ans = ws.cell(r, 3)
            ans.fill = INPUT_FILL
            ans.alignment = MID_LEFT
            ans.font = font(10)
            for c in range(3, self.ncols):
                ws.cell(r, c).border = BOX
                ws.cell(r, c).fill = INPUT_FILL
            g = ws.cell(r, self.ncols, guide)
            g.font = font(9, False, MUTED, italic=True)
            g.alignment = MID_LEFT
            for c in (1, 2, self.ncols):
                ws.cell(r, c).border = ROW_RULE
            if options:
                self._dropdown(f"C{r}", options)
            ws.row_dimensions[r].height = 24
            self.row += 1
        self.form_ranges.append(f"C{start}:C{self.row - 1}")
        self.row += 1

    def table(self, columns: list[tuple], example: list, blank_rows: int = 25,
              prefilled: list[list] | None = None, fixed_cols: int = 0):
        """Data table with a header row, one gold example row and blank input rows.

        columns: (header, options|None, required: bool)
        prefilled: rows whose first `fixed_cols` values we supply (a checklist);
                   the school answers the remaining columns.
        """
        ws = self.ws
        hr = self.row
        for c, (header, _opts, required) in enumerate(columns, start=1):
            cell = ws.cell(hr, c, f"{header} *" if required else header)
            cell.font = font(10, True, WHITE)
            cell.fill = fill(SECONDARY)
            cell.alignment = CENTER
            cell.border = BOX
        ws.row_dimensions[hr].height = 30
        self.row += 1

        if example:
            r = self.row
            for c, v in enumerate(example, start=1):
                cell = ws.cell(r, c, v)
                cell.fill = EXAMPLE_FILL
                cell.font = font(10, False, GOLD_INK, italic=True)
                cell.alignment = MID_LEFT
                cell.border = BOX
            ws.row_dimensions[r].height = 22
            self.row += 1

        first_input_col = fixed_cols + 1
        rows = prefilled if prefilled is not None else [[] for _ in range(blank_rows)]
        start = self.row
        for values in rows:
            r = self.row
            for c in range(1, len(columns) + 1):
                cell = ws.cell(r, c)
                cell.border = BOX
                cell.alignment = MID_LEFT
                if c <= fixed_cols:
                    cell.value = values[c - 1] if c - 1 < len(values) else None
                    cell.font = font(10, c == 2 or fixed_cols == 1, INK)
                    if c == 1 and fixed_cols > 1:
                        cell.font = font(9, False, MUTED)
                        cell.alignment = CENTER
                else:
                    cell.fill = INPUT_FILL
                    cell.font = font(10)
            ws.row_dimensions[r].height = 40 if prefilled else 21
            self.row += 1
        end = self.row - 1

        for c, (_h, opts, _req) in enumerate(columns, start=1):
            if opts and c >= first_input_col:
                col = get_column_letter(c)
                self._dropdown(f"{col}{start}:{col}{end}", opts)

        key_col = get_column_letter(first_input_col)
        if prefilled is not None:
            self.checklists.append((f"{key_col}{start}:{key_col}{end}", len(rows)))
        else:
            self.table_keys.append(f"{key_col}{start}:{key_col}{end}")
        self.row += 1
        return hr

    def note(self, text: str):
        ws, r = self.ws, self.row
        ws.merge_cells(f"A{r}:{get_column_letter(self.ncols)}{r}")
        ws.cell(r, 1, text).font = font(9, False, MUTED, italic=True)
        ws.cell(r, 1).alignment = TOP_LEFT
        ws.row_dimensions[r].height = 30
        self.row += 1

    def finish(self, freeze: str | None = None, header_row: int | None = None):
        ws = self.ws
        if freeze:
            ws.freeze_panes = freeze
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.print_title_rows = "1:2"
        ws.oddFooter.center.text = "Genesis Lumia ERP · School Requirements · Page &P of &N"
        ws.oddFooter.center.size = 8

    # ---- progress (read by Start Here) ----

    def ref(self, rng: str) -> str:
        return f"'{self.name}'!{rng}"

    def progress_formulas(self) -> tuple[str, str]:
        """Returns (progress text formula, status formula) for the Start Here tracker."""
        parts, done_checks, any_checks = [], [], []
        if self.form_ranges:
            # All question blocks on a sheet roll up into one "x of n answers".
            n = sum(_rows_in(rng) for rng in self.form_ranges)
            cnt = "(" + "+".join(f"COUNTA({self.ref(rng)})" for rng in self.form_ranges) + ")"
            parts.append(f'{cnt}&" of {n} answers"')
            done_checks.append(f"{cnt}>={n}")
            any_checks.append(f"{cnt}>0")
        for rng, n in self.checklists:
            cnt = f"COUNTA({self.ref(rng)})"
            parts.append(f'{cnt}&" of {n} decided"')
            done_checks.append(f"{cnt}>={n}")
            any_checks.append(f"{cnt}>0")
        for rng in self.table_keys:
            cnt = f"COUNTA({self.ref(rng)})"
            parts.append(f'{cnt}&IF({cnt}=1," row"," rows")')
            done_checks.append(f"{cnt}>0")
            any_checks.append(f"{cnt}>0")
        text = "=" + '&"  ·  "&'.join(parts)
        status = (
            f'=IF(AND({",".join(done_checks)}),"Complete",'
            f'IF(OR({",".join(any_checks)}),"In progress","Not started"))'
        )
        return text, status

    def _dropdown(self, rng: str, options: list[str]):
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(options) + '"',
            allow_blank=True,
            showErrorMessage=True,
            errorStyle="warning",
            errorTitle="Not in the list",
            error="Pick from the list, or press Yes to keep your own answer.",
        )
        self.ws.add_data_validation(dv)
        dv.add(rng)


def _rows_in(rng: str) -> int:
    a, b = rng.split(":")
    return int("".join(ch for ch in b if ch.isdigit())) - int("".join(ch for ch in a if ch.isdigit())) + 1


# ── Content ──────────────────────────────────────────────────────────────────

def build_profile(wb):
    s = Sheet(wb, "1. School Profile", "1 · School Profile",
              "Who you are, how to reach you, and how your academic year runs. This sets up your school in the ERP.",
              [5, 34, 26, 26, 44])
    s.heading("School identity")
    s.questions([
        ("School name", "Full name as it should appear on receipts and report cards", True, None),
        ("Campus / branch name", "e.g. Main Campus. Leave blank if you have only one campus", False, None),
        ("Board", "Pick one; choose Other and explain in Questions sheet if needed", True,
         ["CBSE", "ICSE", "State Board", "IB", "Cambridge", "Other"]),
        ("Affiliation number", "e.g. 1930001", False, None),
        ("UDISE code", "11-digit code, e.g. 29200800101", False, None),
        ("School type", "Day, residential or both", True, ["Day school", "Residential", "Day + Residential"]),
        ("Students admitted", "Boys, girls or co-ed", True, ["Co-ed", "Boys only", "Girls only"]),
        ("Year established", "e.g. 1962", False, None),
        ("Website", "e.g. www.sainikschool.edu.in", False, None),
    ])
    s.heading("Contact")
    s.questions([
        ("Full address", "Street, area", True, None),
        ("City", "", True, None),
        ("State", "", True, None),
        ("PIN code", "6 digits, e.g. 560001", True, None),
        ("Office phone", "With STD code, e.g. 080-2345 6789", True, None),
        ("Official email", "e.g. office@school.edu.in", True, None),
        ("Principal's name", "e.g. Dr. A. Sharma", True, None),
        ("Principal's email", "", True, None),
        ("Principal's mobile", "10 digits", False, None),
        ("ERP coordinator (our main contact)", "Person who will answer our questions during setup", True, None),
        ("Coordinator mobile & email", "e.g. 98450 12345 · it.coordinator@school.edu.in", True, None),
    ])
    s.heading("Academic year")
    s.questions([
        ("Academic year", "e.g. 2026-27", True, None),
        ("Session start date", "e.g. 01-Apr-2026", True, None),
        ("Session end date", "e.g. 31-Mar-2027", True, None),
        ("Working days", "Which days classes run", True,
         ["Mon–Fri", "Mon–Sat", "Mon–Sat (2nd & 4th Sat off)", "Other"]),
        ("School start time", "e.g. 08:00", True, None),
        ("School end time", "e.g. 14:30", True, None),
    ])
    s.heading("Size (approximate is fine)")
    s.questions([
        ("Total students", "e.g. 850", True, None),
        ("Teaching staff", "e.g. 48", True, None),
        ("Non-teaching staff", "e.g. 30", False, None),
        ("Hostel facility", "", True, YES_NO),
        ("Transport (school buses)", "", True, YES_NO),
    ])
    s.heading("Branding")
    s.questions([
        ("We will send our logo (PNG / SVG)", "High resolution, transparent background if possible", False, YES_NO),
        ("School colours", "e.g. Navy blue and gold", False, None),
        ("Motto / tagline", "Shown on the login page and report cards", False, None),
    ])
    s.finish(freeze="A5")
    return s


def build_modules(wb):
    s = Sheet(wb, "2. Modules", "2 · Modules You Need",
              "Tell us which parts of the ERP you want, how important each is, and anything special about how your school works.",
              [5, 26, 46, 14, 12, 16, 22, 40])
    modules = [
        ("Admissions", "Online parent form, document upload, step-by-step approvals, roll numbers, house allocation, admission letter"),
        ("Student records", "Student master profile, parents, class & section, status, documents"),
        ("Attendance", "Daily roll call, monthly register, absent alerts to parents, attendance %"),
        ("Timetable & substitution", "Auto-generate timetable, teacher workload, cover for absent teachers"),
        ("Academics", "Subjects, syllabus, lesson plans, assignments"),
        ("Exams & report cards", "Exam schedule, marks entry, grades, printable report cards"),
        ("Online tests (CBT)", "Question bank, timed online tests, auto-marking of MCQs"),
        ("Fees & payments", "Fee structure, invoices, payment recording, receipts, dues"),
        ("Leave & gate pass", "Staff and student leave, leave balances, hostel gate passes"),
        ("Hostel", "Hostels, rooms, bed allocation, visitors"),
        ("Discipline", "Incident records, demerits, confidential notes, parent intimation"),
        ("Staff & payroll", "Staff directory, salary structure, monthly payslips, leave deductions"),
        ("Notifications", "SMS / email / WhatsApp to parents and staff, broadcasts"),
        ("Reports & dashboards", "Summary dashboards, exports to Excel/PDF"),
    ]
    s.heading("Modules", "Answer the light-blue columns for every module. Example: Admissions → Yes · High · 15-Mar-2026 · "
                         "Admission in-charge, Principal · “We admit only in Class 6 and 9, after an entrance test.”")
    s.table(
        [("#", None, False), ("Module", None, False), ("What it covers", None, False),
         ("Do you need it?", YES_NO_LATER, True), ("Priority", PRIORITY, False),
         ("Needed by (date)", None, False), ("Main users", None, False),
         ("How it works at your school / special needs", None, False)],
        example=[], prefilled=[[i, m, d] for i, (m, d) in enumerate(modules, start=1)], fixed_cols=3,
    )
    s.finish(freeze="A8")
    return s


def build_classes(wb):
    s = Sheet(wb, "3. Classes", "3 · Classes & Sections",
              "One row per section. This builds your class list, rosters and the timetable grid.",
              [16, 12, 18, 26, 14, 16, 36])
    s.heading("Classes & sections")
    s.table(
        [("Class", None, True), ("Section", None, True),
         ("Stream", ["None", "Science", "Commerce", "Humanities"], False),
         ("Class teacher", None, False), ("Room", None, False),
         ("Students (approx)", None, False), ("Notes", None, False)],
        example=["Class 10", "A", "None", "Mrs. Meera Iyer", "R-204", 40, "Gold row is an example — replace or delete"],
        blank_rows=40,
    )
    s.finish(freeze="A7")
    return s


def build_subjects(wb):
    s = Sheet(wb, "4. Subjects", "4 · Subjects",
              "Every subject taught, the classes it is taught in, and how many periods a week it gets.",
              [26, 12, 18, 20, 14, 12, 34])
    s.heading("Subjects")
    s.table(
        [("Subject name", None, True), ("Short code", None, False), ("Taught in classes", None, True),
         ("Type", ["Theory", "Practical", "Theory + Practical"], False),
         ("Periods / week", None, False), ("Max marks", None, False), ("Notes", None, False)],
        example=["Mathematics", "MAT", "Class 6 – 12", "Theory", 6, 100, "Class 11–12 Science only for Physics"],
        blank_rows=35,
    )
    s.finish(freeze="A7")
    return s


def build_staff(wb):
    s = Sheet(wb, "5. Staff", "5 · Staff & ERP Roles",
              "Everyone who will use the ERP, and what they should be allowed to do. Each person gets their own login.",
              [26, 20, 18, 22, 30, 16, 26, 14, 14])
    s.heading("Staff list", "ERP role decides what each person can see and change. Pick the closest one; explain exceptions in Notes.")
    s.table(
        [("Full name", None, True), ("Designation", None, True), ("Department", None, False),
         ("ERP role", ["Administrator", "Principal", "Vice Principal", "Admission In-charge", "Teacher",
                       "Warden", "Accountant", "Office Staff"], True),
         ("Email (login)", None, True), ("Mobile", None, False), ("Subjects taught", None, False),
         ("Max periods / day", None, False), ("Needs login?", YES_NO, True)],
        example=["Mrs. Meera Iyer", "PGT Mathematics", "Mathematics", "Teacher",
                 "meera.iyer@school.edu.in", "98450 12345", "Mathematics", 6, "Yes"],
        blank_rows=60,
    )
    s.finish(freeze="A8")
    return s


def build_timetable(wb):
    s = Sheet(wb, "6. Timetable", "6 · Timetable Setup",
              "Your bell schedule and the rules the timetable generator must follow.",
              [12, 16, 16, 18, 26, 40])
    s.heading("Rules")
    s.questions([
        ("Periods per day", "Teaching periods only, e.g. 8", True, None),
        ("Period length (minutes)", "e.g. 45", True, None),
        ("Max periods per teacher per day", "e.g. 6", True, None),
        ("Same subject twice in a day?", "Allowed for any class?", False, ["Allowed", "Not allowed", "Only for practicals"]),
        ("Who arranges substitutions?", "e.g. Vice Principal", False, None),
    ])
    s.heading("Bell schedule", "One row per period, break, lunch or assembly — in order.")
    s.table(
        [("Order", None, True), ("Start time", None, True), ("End time", None, True),
         ("Type", ["Class period", "Short break", "Lunch", "Assembly", "Games / Activity"], True),
         ("Name", None, False), ("Notes", None, False)],
        example=[3, "09:30", "09:45", "Short break", "Break", "No classes during breaks"],
        blank_rows=15,
    )
    s.finish(freeze="A5")
    return s


def build_fees(wb):
    s = Sheet(wb, "7. Fees", "7 · Fee Structure",
              "Every fee you charge, for which classes, how often, and when it is due. Amounts in ₹.",
              [20, 22, 15, 18, 22, 22, 32])
    s.heading("Collection")
    s.questions([
        ("Payment modes you accept", "e.g. Cash, UPI, Cheque, NEFT", True, None),
        ("Online payment gateway", "If you already use one, e.g. Razorpay", False, None),
        ("Terms in a year", "", True, ["1 (Annual)", "2", "3", "4 (Quarterly)", "12 (Monthly)"]),
        ("Receipt number format", "e.g. RCPT-2026-0001", False, None),
        ("Sibling / staff-ward concession", "Describe the rule, or 'None'", False, None),
    ])
    s.heading("Fee heads", "One row per fee head per class group.")
    s.table(
        [("Classes", None, True),
         ("Fee head", ["Tuition", "Admission", "Exam", "Hostel", "Mess", "Transport", "Lab",
                       "Activity", "Books & Uniform", "Caution deposit", "Other"], True),
         ("Amount (₹)", None, True),
         ("Frequency", ["One-time", "Annual", "Per term", "Monthly"], True),
         ("Due date", None, True), ("Late fee", None, False), ("Notes", None, False)],
        example=["Class 6 – 8", "Tuition", 12500, "Per term", "10th of the term's first month",
                 "₹50 per day after 15 days", "Refundable: No"],
        blank_rows=30,
    )
    for r in range(s.row - 32, s.row):
        s.ws.cell(r, 3).number_format = '₹#,##0;-₹#,##0;"-"'
    s.finish(freeze="A5")
    return s


def build_admissions(wb):
    s = Sheet(wb, "8. Admissions", "8 · Admissions",
              "How admissions run at your school, and which documents parents must submit.",
              [30, 16, 18, 14, 18, 36])
    s.heading("Process")
    s.questions([
        ("Classes you admit into", "e.g. Class 6 and Class 9 only", True, None),
        ("Admission window", "e.g. January – March", True, None),
        ("Entrance test?", "", True, YES_NO),
        ("Approval steps, in order", "e.g. Office → Accounts → Vice Principal → Principal", True, None),
        ("Online form for parents?", "Parents fill the form themselves via a link", True, YES_NO),
        ("Roll number format", "e.g. 10A-2026-007 (class-section, year, number)", True, None),
        ("House system?", "", True, YES_NO),
        ("House names", "e.g. Red, Blue, Green, Yellow", False, None),
        ("House allocation", "", False, ["Balanced (fill the smallest house)", "Random", "Chosen by school"]),
    ])
    s.heading("Documents parents must submit")
    s.table(
        [("Document", None, True), ("Mandatory?", YES_NO, True),
         ("Accepted format", ["PDF", "JPG / PNG", "PDF or image"], False),
         ("Max size", None, False), ("Original needed at school?", YES_NO, False), ("Notes", None, False)],
        example=["Birth certificate", "Yes", "PDF or image", "2 MB", "Yes", "Municipal certificate only"],
        blank_rows=15,
    )
    s.finish(freeze="A5")
    return s


def build_hostel(wb):
    s = Sheet(wb, "9. Hostel", "9 · Hostel",
              "Skip this sheet if you have no hostel. Otherwise: your hostels, rooms and house rules.",
              [24, 14, 12, 14, 24, 18, 32])
    s.heading("Rules")
    s.questions([
        ("Visitor days & timings", "e.g. Sundays 10:00 – 13:00", False, None),
        ("Gate pass approval", "Who approves a student leaving campus", False,
         ["Warden only", "Warden → Principal", "Principal only"]),
        ("Parent OTP on student return?", "Parent confirms with a one-time code", False, YES_NO),
    ])
    s.heading("Hostels")
    s.table(
        [("Hostel name", None, True), ("For", ["Boys", "Girls", "Mixed"], True),
         ("Rooms", None, True), ("Beds per room", None, True), ("Warden", None, False),
         ("Warden mobile", None, False), ("Notes", None, False)],
        example=["Shivaji House", "Boys", 30, 4, "Mr. Amit Khanna", "98860 22334", "Class 6–8 only"],
        blank_rows=10,
    )
    s.finish(freeze="A5")
    return s


def build_attendance(wb):
    s = Sheet(wb, "10. Attendance & Leave", "10 · Attendance & Leave",
              "How attendance is taken, and the leave types and yearly quotas for staff and students.",
              [26, 14, 16, 14, 16, 22, 30])
    s.heading("Attendance")
    s.questions([
        ("How is attendance taken?", "", True, ["Once a day", "Every period", "Biometric device", "Morning + afternoon"]),
        ("Marked by", "e.g. Class teacher, by 09:00", True, None),
        ("SMS parents when a child is absent?", "", True, YES_NO),
        ("Minimum attendance required (%)", "e.g. 75", False, None),
        ("Biometric device brand (if any)", "e.g. eSSL, ZKTeco, CAMS", False, None),
    ])
    s.heading("Leave types")
    s.table(
        [("Leave type", None, True), ("For", ["Staff", "Students", "Both"], True),
         ("Days per year", None, True), ("Paid?", ["Paid", "Unpaid"], True),
         ("Carry forward?", YES_NO, False), ("Approved by", None, True), ("Notes", None, False)],
        example=["Casual Leave", "Staff", 12, "Paid", "No", "Principal", "Max 3 days at a time"],
        blank_rows=12,
    )
    s.finish(freeze="A5")
    return s


def build_exams(wb):
    s = Sheet(wb, "11. Exams", "11 · Exams",
              "Every exam in the year, the classes that sit it, and how much it counts towards the final result.",
              [28, 18, 16, 12, 14, 16, 30])
    s.heading("Rules")
    s.questions([
        ("Result shown as", "", True, ["Marks", "Grades", "Marks + Grades"]),
        ("Pass mark (%)", "e.g. 33", True, None),
        ("We will send a sample report card", "So the printed report card matches yours", False, YES_NO),
    ])
    s.heading("Exams in the year")
    s.table(
        [("Exam name", None, True), ("Classes", None, True), ("Month", None, True),
         ("Max marks", None, True), ("Weightage (%)", None, False),
         ("Type", ["Written", "Practical", "Online (CBT)", "Mixed"], False), ("Notes", None, False)],
        example=["Half-Yearly Examination", "Class 6 – 12", "September", 80, 30, "Written", "+20 internal assessment"],
        blank_rows=12,
    )
    s.finish(freeze="A5")
    return s


def build_grading(wb):
    s = Sheet(wb, "12. Grading", "12 · Grading Scale",
              "Your grade bands. Leave blank if you report marks only.",
              [12, 14, 14, 16, 34])
    s.heading("Grade bands", "Highest grade first. Percentages are inclusive.")
    s.table(
        [("Grade", None, True), ("From (%)", None, True), ("To (%)", None, True),
         ("Grade point", None, False), ("Remark on report card", None, False)],
        example=["A1", 91, 100, 10, "Outstanding"],
        blank_rows=12,
    )
    s.finish(freeze="A8")
    return s


def build_notifications(wb):
    s = Sheet(wb, "13. Notifications", "13 · Notifications",
              "Which messages go out automatically, to whom, and by which channel.",
              [5, 34, 18, 18, 14, 34])
    s.heading("Channels")
    s.questions([
        ("SMS provider (if any)", "e.g. MSG91, Textlocal — or None", False, None),
        ("SMS sender ID", "6 letters, e.g. SAINIK", False, None),
        ("Email sender address", "e.g. noreply@school.edu.in", False, None),
        ("WhatsApp Business number", "If you have one", False, None),
    ])
    s.heading("Automatic messages", "Decide each one. Channel can be SMS, Email, WhatsApp or several.")
    events = [
        "Student absent today", "Fee due reminder", "Fee payment receipt", "Exam results published",
        "Leave approved / rejected", "Gate pass issued / student returned", "Admission status update",
        "Holiday / general announcement", "Discipline incident (no confidential details)",
        "New assignment posted", "Payslip ready (staff)",
    ]
    s.table(
        [("#", None, False), ("Message", None, False), ("Send to", ["Parents", "Staff", "Students", "Parents + Staff"], False),
         ("Channel", ["SMS", "Email", "WhatsApp", "SMS + Email", "All"], False), ("Needed?", YES_NO, True),
         ("Notes", None, False)],
        example=[], prefilled=[[i, e] for i, e in enumerate(events, start=1)], fixed_cols=2,
    )
    s.finish(freeze="A5")
    return s


def build_existing_data(wb):
    s = Sheet(wb, "14. Existing Data", "14 · Existing Data",
              "What records you already have, so we can import them instead of typing them in again.",
              [5, 30, 14, 24, 16, 36])
    s.heading("Data you can send us", "We will share a simple import format for each 'Yes'.")
    items = [
        "Student list (name, class, parents, contacts)", "Parent / guardian contacts", "Staff list",
        "Fee dues & payments this year", "Past attendance", "Past exam marks", "Current timetable",
        "Hostel room allotments", "Salary details (for payroll)", "Library books",
    ]
    s.table(
        [("#", None, False), ("Data", None, False), ("Available?", YES_NO, True),
         ("Kept in", ["Excel / Google Sheets", "Paper register", "Other software", "Not kept"], False),
         ("Approx. records", None, False), ("Notes (software name, years covered…)", None, False)],
        example=[], prefilled=[[i, d] for i, d in enumerate(items, start=1)], fixed_cols=2,
    )
    s.finish(freeze="A5")
    return s


def build_reports(wb):
    s = Sheet(wb, "15. Reports", "15 · Reports You Need",
              "Reports, registers and certificates your office produces today — we will match them.",
              [30, 24, 16, 18, 36])
    s.heading("Reports")
    s.table(
        [("Report / certificate", None, True), ("Who uses it", None, False),
         ("How often", ["Daily", "Weekly", "Monthly", "Per term", "Yearly", "On request"], False),
         ("Format", ["Printed", "PDF", "Excel", "On screen"], False),
         ("Notes (please attach a sample if possible)", None, False)],
        example=["Monthly fee defaulters list", "Accountant, Principal", "Monthly", "Excel", "Sorted by class"],
        blank_rows=20,
    )
    s.finish(freeze="A7")
    return s


def build_signoff(wb):
    s = Sheet(wb, "16. Questions & Sign-off", "16 · Questions & Sign-off",
              "Anything we did not ask, any doubts — and who filled in this workbook.",
              [5, 52, 22, 14, 30])
    s.heading("Your questions & special requirements")
    s.table(
        [("#", None, False), ("Question or requirement", None, True), ("Raised by", None, False),
         ("Priority", PRIORITY, False), ("Notes", None, False)],
        example=[1, "Can parents see their child's attendance on their phone?", "Principal", "Medium", ""],
        blank_rows=15, fixed_cols=1, prefilled=None,
    )
    s.heading("Sign-off")
    s.questions([
        ("Filled in by", "Name", True, None),
        ("Designation", "", True, None),
        ("Date", "e.g. 20-Oct-2026", True, None),
        ("Approved by (Principal)", "Name", False, None),
    ])
    s.finish(freeze="A5")
    return s


# ── Start Here ───────────────────────────────────────────────────────────────

def build_start(ws, sections: list[tuple[Sheet, str]], profile: Sheet):
    ws.title = "Start Here"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    widths = [4, 5, 28, 46, 30, 15, 12, 4]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Hero band.
    for r in range(1, 5):
        for c in range(1, 9):
            ws.cell(r, c).fill = fill(NAVY if r < 4 else DEEP)
    ws.merge_cells("B2:G2")
    ws["B2"] = "Genesis Lumia ERP"
    ws["B2"].font = font(22, True, WHITE)
    ws["B2"].alignment = Alignment(vertical="center")
    ws.merge_cells("B3:G3")
    ws["B3"] = "School Requirements Workbook"
    ws["B3"].font = font(13, False, "B4CBE6")
    ws["B3"].alignment = Alignment(vertical="top")
    ws.merge_cells("B4:G4")
    ws["B4"] = "Please fill this in and send it back — it is everything we need to set up the ERP for your school."
    ws["B4"].font = font(10, False, "D6E4F4", italic=True)
    ws["B4"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 14
    ws.row_dimensions[2].height = 36
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 26

    # School + return details.
    r = 6
    ws.merge_cells(f"B{r}:D{r}")
    ws.cell(r, 2, "SCHOOL").font = font(8, True, MUTED)
    ws.merge_cells(f"E{r}:G{r}")
    ws.cell(r, 5, "PLEASE RETURN THE COMPLETED FILE TO").font = font(8, True, MUTED)
    ws.cell(r, 5).alignment = Alignment(horizontal="right")
    ws.merge_cells(f"B{r + 1}:D{r + 1}")
    name_ref = profile.ref(profile.form_ranges[0].split(":")[0])
    ws.cell(r + 1, 2).value = f'=IF({name_ref}="","(fills in automatically from Sheet 1)",{name_ref})'
    ws.cell(r + 1, 2).font = font(13, True, NAVY)
    ws.merge_cells(f"E{r + 1}:G{r + 1}")
    ws.cell(r + 1, 5, RETURN_TO).font = font(10, True, SECONDARY)
    ws.cell(r + 1, 5).alignment = Alignment(horizontal="right")
    ws.row_dimensions[r + 1].height = 24

    # How to fill.
    r = 9
    ws.cell(r, 2, "HOW TO FILL THIS IN").font = font(10, True, NAVY)
    for c in range(2, 8):
        ws.cell(r, c).border = Border(bottom=Side(style="medium", color=NAVY))
    steps = [
        "Work through the numbered sheets (tabs along the bottom). Each one takes 5–15 minutes.",
        "Type only in the light-blue cells. Where a cell has a ▾ arrow, pick from the list.",
        "The gold row on each sheet is an example of the format — overwrite it or leave it.",
        "Fields marked * are required. If something does not apply to you, write 'NA'.",
        "Unsure about anything? Note it in Sheet 16 — we will discuss it with you.",
        "Save the file and send it back. Sending it in parts is fine — the tracker below shows what's left.",
    ]
    for i, step in enumerate(steps, start=1):
        rr = r + i
        ws.cell(rr, 2, i).font = font(10, True, WHITE)
        ws.cell(rr, 2).fill = fill(SECONDARY)
        ws.cell(rr, 2).alignment = CENTER
        ws.merge_cells(f"C{rr}:G{rr}")
        ws.cell(rr, 3, step).font = font(10)
        ws.cell(rr, 3).alignment = MID_LEFT
        ws.row_dimensions[rr].height = 22

    # Legend.
    r = r + len(steps) + 2
    ws.cell(r, 2, "COLOUR KEY").font = font(10, True, NAVY)
    for c in range(2, 8):
        ws.cell(r, c).border = Border(bottom=Side(style="medium", color=NAVY))
    legend = [
        (PALE, "", "Light blue — please fill in"),
        (GOLD_SOFT, "e.g.", "Gold — an example of the expected format"),
        (WHITE, "", "White — questions and guidance from us (no need to change)"),
    ]
    for i, (hex_, txt, label) in enumerate(legend, start=1):
        rr = r + i
        sw = ws.cell(rr, 2, txt)
        sw.fill = fill(hex_)
        sw.border = BOX
        sw.font = font(8, False, GOLD_INK, italic=True)
        sw.alignment = CENTER
        ws.merge_cells(f"C{rr}:G{rr}")
        ws.cell(rr, 3, label).font = font(10)
        ws.cell(rr, 3).alignment = MID_LEFT
        ws.row_dimensions[rr].height = 20

    # Progress tracker.
    r = r + len(legend) + 2
    ws.cell(r, 2, "YOUR PROGRESS").font = font(10, True, NAVY)
    for c in range(2, 8):
        ws.cell(r, c).border = Border(bottom=Side(style="medium", color=NAVY))
    hr = r + 1
    for c, h in enumerate(["#", "Section", "What we need", "Filled so far", "Status", "Go"], start=2):
        cell = ws.cell(hr, c, h)
        cell.font = font(10, True, WHITE)
        cell.fill = fill(SECONDARY)
        cell.alignment = CENTER
        cell.border = BOX
    ws.row_dimensions[hr].height = 24
    first = hr + 1
    for i, (sheet, need) in enumerate(sections, start=1):
        rr = hr + i
        text_f, status_f = sheet.progress_formulas()
        values = [i, sheet.name.split(". ", 1)[1], need, text_f, status_f,
                  f'=HYPERLINK("#\'{sheet.name}\'!A1","Open →")']
        for c, v in enumerate(values, start=2):
            cell = ws.cell(rr, c, v)
            cell.border = BOX
            cell.alignment = MID_LEFT
            cell.font = font(10)
        ws.cell(rr, 2).alignment = CENTER
        ws.cell(rr, 2).font = font(9, False, MUTED)
        ws.cell(rr, 3).font = font(10, True, INK)
        ws.cell(rr, 4).font = font(9, False, MUTED)
        ws.cell(rr, 5).font = font(9, False, INK)
        ws.cell(rr, 6).alignment = CENTER
        ws.cell(rr, 6).font = font(9, True, INK)
        ws.cell(rr, 7).font = font(9, True, SECONDARY, underline="single")
        ws.cell(rr, 7).alignment = CENTER
        ws.row_dimensions[rr].height = 22
    last = hr + len(sections)

    status_rng = f"F{first}:F{last}"
    for word, bg, fg in (("Complete", OK_SOFT, OK_INK), ("In progress", WARN_SOFT, WARN_INK),
                         ("Not started", SLATE_SOFT, MUTED)):
        ws.conditional_formatting.add(
            status_rng,
            FormulaRule(formula=[f'F{first}="{word}"'], fill=fill(bg), font=Font(name=FONT, bold=True, color=fg)),
        )

    rr = last + 2
    ws.merge_cells(f"B{rr}:E{rr}")
    ws.cell(rr, 2).value = "Sections complete"
    ws.cell(rr, 2).font = font(11, True, NAVY)
    ws.cell(rr, 2).alignment = Alignment(horizontal="right", vertical="center")
    ws.merge_cells(f"F{rr}:G{rr}")
    ws.cell(rr, 6).value = f'=COUNTIF({status_rng},"Complete")&" of {len(sections)}"'
    ws.cell(rr, 6).font = font(12, True, WHITE)
    ws.cell(rr, 6).fill = fill(NAVY)
    ws.cell(rr, 6).alignment = CENTER
    ws.row_dimensions[rr].height = 26

    ws.merge_cells(f"B{rr + 2}:G{rr + 2}")
    ws.cell(rr + 2, 2).value = ("Your answers are used only to configure your school's ERP. "
                                "Questions? Contact " + RETURN_TO.replace("  ·  ", " at ") + ".")
    ws.cell(rr + 2, 2).font = font(9, False, MUTED, italic=True)
    ws.cell(rr + 2, 2).alignment = MID_LEFT

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    wb = Workbook()
    start = wb.active
    profile = build_profile(wb)
    sections = [
        (profile, "Name, board, contacts, academic year, size"),
        (build_modules(wb), "Which ERP modules you want and their priority"),
        (build_classes(wb), "Every class and section, with class teacher"),
        (build_subjects(wb), "Subjects, classes taught in, periods per week"),
        (build_staff(wb), "Everyone who needs a login, and their role"),
        (build_timetable(wb), "Bell schedule and timetable rules"),
        (build_fees(wb), "Fee heads, amounts, terms and due dates"),
        (build_admissions(wb), "Admission process and required documents"),
        (build_hostel(wb), "Hostels, rooms, visitor & gate-pass rules"),
        (build_attendance(wb), "Attendance method, leave types and quotas"),
        (build_exams(wb), "Exams in the year and their weightage"),
        (build_grading(wb), "Grade bands (if you use grades)"),
        (build_notifications(wb), "Automatic SMS / email / WhatsApp messages"),
        (build_existing_data(wb), "Records we can import for you"),
        (build_reports(wb), "Reports and certificates you use today"),
        (build_signoff(wb), "Open questions and who filled this in"),
    ]
    build_start(start, sections, profile)
    wb.active = 0
    wb.properties.title = "Genesis Lumia ERP — School Requirements Workbook"
    wb.properties.creator = "Genesis Lumia ERP"
    wb.save(OUT)
    print(f"Wrote {OUT} with {len(wb.sheetnames)} sheets")


if __name__ == "__main__":
    main()
